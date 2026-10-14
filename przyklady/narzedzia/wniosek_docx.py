#!/usr/bin/env python3
"""Składa wniosek.docx z wniosek.md przykładu (python-docx).

Użycie:
  python3 przyklady/narzedzia/wniosek_docx.py przyklady/01-fryzjerka-mobilna-turek

Strona tytułowa: nazwa urzędu, naboru, instrukcja „co wpisać samemu” (blok '@@ docx_samemu'
i '@@ docx_tytul' z narzedzia/notatki/<katalog>.md). Dalej treść wniosek.md w kolejności pól
urzędu: nagłówki, nazwy pól (wyróżnione), akapity, listy i tabele (nagłówek tabeli powtarza się
na kolejnych stronach). Tekst ma pozostać edytowalny – klient kopiuje go do formularza urzędu.
"""
import os, re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))

def notes(folder):
    p = os.path.join(HERE, "notatki", os.path.basename(os.path.normpath(folder)) + ".md")
    out, cur = {}, None
    for line in open(p, encoding="utf-8"):
        m = re.match(r"^@@ (\w+)", line)
        if m:
            cur = m.group(1); out[cur] = []
            continue
        if cur:
            out[cur].append(line.rstrip("\n"))
    return {k: [l for l in v] for k, v in out.items()}

def shade(cell_or_par, fill="E7E6E6"):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(cell_or_par, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pr.append(shd)

def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trPr.append(el)

def add_runs(par, text, size=None, italic=False):
    # **pogrubienie** w treści
    parts = re.split(r"(\*\*.+?\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        t = part[2:-2] if bold else part
        t = t.replace("`", "")
        r = par.add_run(t)
        r.bold = bold
        r.italic = italic
        if size:
            r.font.size = Pt(size)
    return par

def table_from(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    ncol = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        row = t.add_row()
        for j in range(ncol):
            txt = r[j] if j < len(r) else ""
            c = row.cells[j]
            c.text = ""
            p = c.paragraphs[0]
            add_runs(p, txt, size=9)
            if i == 0:
                for run in p.runs:
                    run.bold = True
                shade(c, "D9D9D9")
        if i == 0:
            repeat_header(row)
    doc.add_paragraph()
    return t

def title_page(doc, N):
    for line in N.get("docx_tytul", []):
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(line[2:]); r.bold = True; r.font.size = Pt(18)
        else:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs(p, line, size=11)
    doc.add_paragraph()
    h = doc.add_paragraph(); r = h.add_run("Zanim złożysz wniosek – co musisz wpisać i zrobić SAM(A)"); r.bold = True; r.font.size = Pt(13)
    shade(h, "FFF2CC")
    for line in N.get("docx_samemu", []):
        if not line.strip():
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, line[2:], size=10)
        else:
            p = doc.add_paragraph(); add_runs(p, line, size=10)
    p = doc.add_paragraph()
    add_runs(p, "Jak korzystać z tego pliku: każda część zaczyna się od szarego paska z nazwą pola z formularza urzędu. "
                "Tekst pod paskiem przepisz lub wklej do tego pola. Nie zmieniaj liczb – zgadzają się z arkuszem kalkulacyjnym "
                "dołączonym do zlecenia. Jeśli zmienisz zakup lub cenę, najpierw zmień arkusz.", size=10, italic=True)
    p.add_run().add_break(WD_BREAK.PAGE)

def main(folder):
    N = notes(folder)
    md = open(os.path.join(folder, "wniosek.md"), encoding="utf-8").read().splitlines()
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    title_page(doc, N)
    i, in_map = 0, False
    buf_table = []
    def flush_table():
        nonlocal buf_table
        if buf_table:
            table_from(doc, buf_table)
            buf_table = []
    while i < len(md):
        line = md[i]
        if line.startswith("|"):
            buf_table.append(line); i += 1; continue
        flush_table()
        if line.startswith("# "):
            doc.add_heading(line[2:], level=0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("> "):
            p = doc.add_paragraph(); add_runs(p, line[2:], size=9, italic=True); shade(p, "F2F2F2")
        elif line.startswith("*Pol") and line.endswith("*"):
            p = doc.add_paragraph(); add_runs(p, "Pole formularza: " + line.strip("*").split(":", 1)[1].strip(), size=10)
            for r in p.runs:
                r.bold = True
            shade(p, "D9D9D9")
        elif line.startswith("– ") or line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, line[2:])
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(); add_runs(p, line)  # numer zostaje w tekście (bez autonumeracji Worda)
        elif line.strip() == "---":
            pass
        elif line.strip():
            p = doc.add_paragraph(); add_runs(p, line)
        i += 1
    flush_table()
    out = os.path.join(folder, "wniosek.docx")
    doc.save(out)
    print("zapisano", out)

if __name__ == "__main__":
    main(sys.argv[1])
