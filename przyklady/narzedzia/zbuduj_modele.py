#!/usr/bin/env python3
"""Buduje model.xlsx przykładów z szablonu model-finansowy/model-finansowy-pup.xlsx.

Użycie (z katalogu repozytorium):
  python3 przyklady/narzedzia/zbuduj_modele.py            # wszystkie przykłady
  python3 przyklady/narzedzia/zbuduj_modele.py mechanik   # tylko przykłady, których nazwa zawiera słowo

Dla każdego przykładu: kopiuje szablon, wpisuje dane wejściowe (żółte komórki: Zalozenia,
Kosztorys, Przychody, Koszty), dobiera cenę pozycji „towar”, żeby suma z dotacji była równa
kwocie wniosku co do grosza, przelicza plik przez LibreOffice
(/usr/bin/soffice --headless --convert-to xlsx) i zapisuje przyklady/<katalog>/model.xlsx.
Nowy przykład = nowy wpis w słowniku EXAMPLES.
"""
import datetime, json, os, shutil, subprocess, sys, tempfile
import openpyxl

REPO = os.path.abspath(os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", ".."))
TPL = f"{REPO}/model-finansowy/model-finansowy-pup.xlsx"
SCR = tempfile.mkdtemp(prefix="modele_")

def ceneo(q):
    return "https://www.ceneo.pl/;szukaj-" + q.replace(" ", "+")

EXAMPLES = {
 "01-fryzjerka-mobilna-turek": dict(
  zal=dict(B4="Fryzjerka mobilna – Turek (DANE FIKCYJNE – przykład)", B5=15000, B6=38822, B7=4000,
           B8=datetime.datetime(2026, 12, 1), B9="ryczałt", B10=0.085, B11="TAK", B12=6, B13="TAK",
           B14=24, B15="NIE", B16=3000, B17="TAK", B18="NIE", B19=0.08, B20=21, B21=6, B22=0.3),
  target_dotacja=15000.00, adjust_row="towar",
  items=[  # nazwa, kategoria, nowy/używany, ilość, cena netto, VAT, link, finansowanie
   ("Przenośna umywalka fryzjerska z kranem, zbiornikiem i regulacją wysokości", "sprzęt", "nowy", 1, 640, 0.23, "https://www.kaufland.pl/product/463416684/", "dotacja"),
   ("Nadmuchiwana wanienka do mycia głowy w łóżku", "sprzęt", "nowy", 2, 65, 0.23, ceneo("wanienka do mycia glowy w lozku"), "dotacja"),
   ("Maszynka bezprzewodowa Wahl Magic Clip Cordless", "sprzęt", "nowy", 1, 520, 0.23, "https://www.okazje.info.pl/szukaj/wahl-maszynka-do-strzyzenia-wlosow-magic-clip-5-star-cordless/ceny.html", "dotacja"),
   ("Trymer bezprzewodowy Wahl Detailer Li", "sprzęt", "nowy", 1, 430, 0.23, ceneo("wahl detailer li"), "dotacja"),
   ("Suszarka Parlux Advance Light 2200 W", "sprzęt", "nowy", 1, 490, 0.23, "https://www.notino.pl/parlux/advance-light-suszarka-do-wlosow/", "dotacja"),
   ("Prostownica BaByliss PRO EP Technology 5.0 (24 mm)", "sprzęt", "nowy", 1, 300, 0.23, "https://www.notino.pl/babyliss-pro/straighteners-ep-technology-50-2072e-prostownica-do-wlosow/", "dotacja"),
   ("Nożyczki Jaguar Pre Style Relax 5,5\" + degażówki Jaguar Ergo Pre Style 5,5\"", "sprzęt", "nowy", 1, 620, 0.23, "https://www.kaufland.pl/product/306979817/", "dotacja"),
   ("Sterylizator UV-C i myjka ultradźwiękowa do narzędzi", "sprzęt", "nowy", 1, 450, 0.23, ceneo("sterylizator uv fryzjerski"), "dotacja"),
   ("Walizka fryzjerska na kółkach z organizerami", "sprzęt", "nowy", 1, 450, 0.23, ceneo("walizka fryzjerska na kolkach"), "dotacja"),
   ("Krzesło fryzjerskie składane z regulacją wysokości", "sprzęt", "nowy", 1, 520, 0.23, ceneo("krzeslo fryzjerskie skladane"), "dotacja"),
   ("Lampa LED akumulatorowa na statywie", "sprzęt", "nowy", 1, 280, 0.23, ceneo("lampa led akumulatorowa statyw"), "dotacja"),
   ("Odkurzacz ręczny akumulatorowy (sprzątanie włosów u klienta)", "sprzęt", "nowy", 1, 400, 0.23, ceneo("odkurzacz reczny akumulatorowy"), "dotacja"),
   ("Tablet 10\" do kalendarza wizyt i zdjęć fryzur", "sprzęt", "nowy", 1, 900, 0.23, ceneo("tablet 10 cali"), "dotacja"),
   ("Peleryny, ręczniki (30 szt.), grzebienie, szczotki, klipsy, miski, pędzle", "inne", "nowy", 1, 800, 0.23, ceneo("peleryna fryzjerska"), "dotacja"),
   ("Farby, oksydanty, rozjaśniacze, szampony – pierwszy zapas na 2–3 mies.", "towar", "nowy", 1, 2600, 0.23, "[oferta hurtowni fryzjerskiej – do dołączenia przez klientkę]", "dotacja"),
   ("Oklejenie samochodu folią z logo i telefonem", "reklama", "nowy", 1, 1100, 0.23, "[oferta lokalnej firmy reklamowej – do dołączenia]", "dotacja"),
   ("Logo i strona www z kalendarzem rezerwacji", "reklama", "nowy", 1, 1500, 0.23, "[oferta wykonawcy strony – do dołączenia]", "dotacja"),
   ("Kasa fiskalna online (wyłączona z dofinansowania w PUP Turek)", "sprzęt", "nowy", 1, 1220, 0.23, ceneo("kasa fiskalna online"), "wkład własny"),
   ("Ulotki, wizytówki, karnety stałej klientki", "reklama", "nowy", 1, 300, 0.23, "[oferta drukarni – do dołączenia]", "wkład własny"),
   ("Tunika robocza – 2 szt.", "inne", "nowy", 1, 150, 0.23, ceneo("tunika fryzjerska"), "wkład własny"),
  ],
  limits=dict(sprzęt=1, towar=1, reklama=1, remont=0, inne=1),
  services=[  # nazwa, cena, liczba przy pełnym obłożeniu, czas (h)
   ("Strzyżenie damskie z myciem i modelowaniem", 70, 35, 0.75),
   ("Strzyżenie męskie / dziecięce", 40, 25, 0.5),
   ("Strzyżenie seniora w domu (także osoby leżące)", 45, 25, 0.6),
   ("Koloryzacja z farbą (włosy krótkie i średnie)", 180, 12, 2.0),
   ("Upięcie okolicznościowe / ślubne", 180, 3, 1.5),
  ],
  season=[0.85, 0.85, 0.95, 1.05, 1.10, 1.15, 1.00, 0.95, 1.05, 0.95, 1.00, 1.15],
  rozruch=[0.5, 0.7, 0.9] + [1]*9,
  koszty=[("Czynsz / najem lokalu (nie dotyczy – praca mobilna)", 0), ("Media (nie dotyczy)", 0),
          ("Telefon i internet", 70), ("Księgowość (biuro rachunkowe)", 200), ("Ubezpieczenie OC działalności", 40),
          ("Paliwo i eksploatacja auta", 650), ("Reklama bieżąca (Facebook, Google)", 150),
          ("Inne (abonament Booksy, konto firmowe, serwis kasy)", 150)],
  materialy=0.10,
 ),
 "02-mechanik-grodzisk-mazowiecki": dict(
  zal=dict(B4="Warsztat mechaniczny – Kałęczyn, gm. Grodzisk Maz. (DANE FIKCYJNE – przykład)", B5=39400, B6=55398.78, B7=8000,
           B8=datetime.datetime(2027, 2, 1), B9="skala", B10=0.085, B11="TAK", B12=6, B13="TAK",
           B14=24, B15="NIE", B16=3000, B17="TAK", B18="NIE", B19=0.23, B20=21, B21=7, B22=0.3),
  target_dotacja=39400.00, adjust_row="towar",
  items=[
   ("Podnośnik dwukolumnowy 4 t elektrohydrauliczny z montażem i dokumentacją UDT", "sprzęt", "nowy", 1, 7725, 0.23, "https://www.otomoto.pl/oferta/podnosnik-dwukolumnowy-euro-trade-sae-40-4000-kg-automatyczny-elektrohydrauliczny-ce-udt-ID6IaxBY.html", "dotacja"),
   ("Tester diagnostyczny wielomarkowy (klasa Launch X-431 / Autel MaxiSys) z aktualizacją 12 mies.", "sprzęt", "nowy", 1, 4800, 0.23, ceneo("tester diagnostyczny launch x431"), "dotacja"),
   ("Kompresor tłokowy 270 l, 400 V", "sprzęt", "nowy", 1, 2600, 0.23, ceneo("kompresor 270l 400v"), "dotacja"),
   ("Klucz udarowy pneumatyczny 1/2\" + akumulatorowy, z kompletem nasadek udarowych", "sprzęt", "nowy", 1, 1400, 0.23, ceneo("klucz udarowy 1/2"), "dotacja"),
   ("Prasa hydrauliczna warsztatowa 20 t", "sprzęt", "nowy", 1, 1300, 0.23, ceneo("prasa hydrauliczna 20t"), "dotacja"),
   ("Zestaw blokad rozrządu (VAG, PSA, Renault, Ford, Opel – benzyna i diesel)", "sprzęt", "nowy", 1, 900, 0.23, ceneo("blokada rozrzadu zestaw"), "dotacja"),
   ("Wysysarka/zlewarka do oleju pneumatyczna 80 l", "sprzęt", "nowy", 1, 750, 0.23, ceneo("zlewarka oleju pneumatyczna 80l"), "dotacja"),
   ("Podnośnik żaba 3 t niskoprofilowy + 2 pary stojaków", "sprzęt", "nowy", 1, 700, 0.23, ceneo("podnosnik zaba 3t niski"), "dotacja"),
   ("Prostownik z funkcją rozruchu i zasilacz stabilizujący do kodowania", "sprzęt", "nowy", 1, 800, 0.23, ceneo("zasilacz do kodowania samochodu"), "dotacja"),
   ("Myjka do części z pompą", "sprzęt", "nowy", 1, 650, 0.23, ceneo("myjka do czesci z pompa"), "dotacja"),
   ("Urządzenie do wymiany płynu hamulcowego", "sprzęt", "nowy", 1, 900, 0.23, ceneo("urzadzenie do wymiany plynu hamulcowego"), "dotacja"),
   ("Laptop do diagnostyki, katalogów części i kosztorysów", "sprzęt", "nowy", 1, 2400, 0.23, ceneo("laptop 15,6 16gb"), "dotacja"),
   ("Stół warsztatowy z imadłem + szafka narzędziowa", "sprzęt", "nowy", 1, 1600, 0.23, ceneo("stol warsztatowy z imadlem"), "dotacja"),
   ("Oleje silnikowe (2 × 60 l), filtry, płyny – pierwszy zapas", "towar", "nowy", 1, 3300, 0.23, "[oferta hurtowni motoryzacyjnej – do dołączenia]", "dotacja"),
   ("Kaseton z nazwą na halę + baner przy drodze", "reklama", "nowy", 1, 1000, 0.23, "[oferta firmy reklamowej – do dołączenia]", "dotacja"),
   ("Strona www z mapą dojazdu i formularzem zgłoszenia", "reklama", "nowy", 1, 1200, 0.23, "[oferta wykonawcy – do dołączenia]", "dotacja"),
   ("Kasa fiskalna online", "sprzęt", "nowy", 1, 1200, 0.23, ceneo("kasa fiskalna online"), "wkład własny"),
   ("Gaśnice, apteczka, sorbent, wanny wychwytowe (BHP i ochrona środowiska)", "inne", "nowy", 1, 400, 0.23, ceneo("sorbent do oleju"), "wkład własny"),
   ("Ubranie robocze i obuwie ochronne", "inne", "nowy", 1, 300, 0.23, ceneo("ubranie robocze mechanika"), "wkład własny"),
  ],
  limits=dict(sprzęt=1, towar=0.30, reklama=0.15, remont=0, inne=1),
  services=[
   ("Serwis olejowy – robocizna (olej + filtry)", 70, 40, 0.5),
   ("Hamulce i zawieszenie – robocizna (średnio na zlecenie)", 230, 22, 1.5),
   ("Diagnostyka komputerowa i elektryka", 120, 18, 0.75),
   ("Rozrząd / sprzęgło – robocizna (średnio)", 700, 4, 5.0),
   ("Części i oleje odsprzedane klientom (średnio na zlecenie)", 110, 36, 0.0),
  ],
  season=[0.90, 0.90, 1.05, 1.10, 1.05, 1.00, 0.90, 0.90, 1.05, 1.10, 1.10, 0.95],
  rozruch=[0.4, 0.6, 0.8] + [1]*9,
  koszty=[("Czynsz – najem hali w Kałęczynie", 2200), ("Media (prąd 400 V, woda, ogrzewanie)", 450),
          ("Telefon i internet", 100), ("Księgowość (biuro rachunkowe, KPiR)", 350), ("Ubezpieczenie OC działalności", 80),
          ("Paliwo (dojazdy po części, do klientów)", 300), ("Reklama bieżąca", 200),
          ("Inne (licencja testera od 2. roku – rezerwa, odbiór olejów i odpadów, konto)", 300)],
  materialy=0.26,
 ),
 "03-firma-sprzatajaca-kolbuszowa": dict(
  zal=dict(B4="Firma sprzątająca – Kolbuszowa (DANE FIKCYJNE – przykład)", B5=24850, B6=40000, B7=8000,
           B8=datetime.datetime(2027, 3, 1), B9="ryczałt", B10=0.085, B11="TAK", B12=6, B13="TAK",
           B14=24, B15="NIE", B16=3000, B17="TAK", B18="NIE", B19=0.08, B20=21, B21=7, B22=0.3),
  target_dotacja=24850.00, adjust_row="towar",
  items=[
   ("Odkurzacz piorący Kärcher Puzzi 10/1", "sprzęt", "nowy", 1, 3900, 0.23, ceneo("karcher puzzi 10/1"), "dotacja"),
   ("Odkurzacz plecakowy Kärcher BV 5/1 Bp (schody, klatki)", "sprzęt", "nowy", 1, 1550, 0.23, ceneo("karcher bv 5/1"), "dotacja"),
   ("Odkurzacz na sucho Kärcher T 10/1 Adv HEPA", "sprzęt", "nowy", 1, 1100, 0.23, ceneo("karcher t 10/1 adv"), "dotacja"),
   ("Szorowarka kompaktowa Kärcher BR 30/4 C", "sprzęt", "nowy", 1, 6900, 0.23, ceneo("karcher br 30/4 c"), "dotacja"),
   ("Myjka do okien Kärcher WV 6 Plus", "sprzęt", "nowy", 1, 400, 0.23, ceneo("karcher wv 6 plus"), "dotacja"),
   ("Wózek serwisowy z prasą do mopów i workiem na odpady", "sprzęt", "nowy", 1, 1100, 0.23, ceneo("wozek serwisowy do sprzatania z prasa"), "dotacja"),
   ("Mopy płaskie i ściereczki z mikrofibry w 4 kolorach (system stref), uchwyty", "inne", "nowy", 1, 650, 0.23, ceneo("mop plaski mikrofibra zestaw"), "dotacja"),
   ("Środki czystości – koncentraty z dozownikami (Ecolabel), pierwszy zapas na 3 mies.", "towar", "nowy", 1, 1800, 0.23, "[oferta hurtowni chemii profesjonalnej – do dołączenia]", "dotacja"),
   ("Oklejenie samochodu folią z nazwą i telefonem", "reklama", "nowy", 1, 1300, 0.23, "[oferta firmy reklamowej – do dołączenia]", "dotacja"),
   ("Strona www – wizytówka z cennikiem i formularzem", "reklama", "nowy", 1, 1500, 0.23, "[oferta wykonawcy – do dołączenia]", "dotacja"),
   ("Parownica Kärcher SC 5 EasyFix (sprzątanie bez chemii)", "sprzęt", "nowy", 1, 1300, 0.23, ceneo("karcher sc 5 easyfix"), "wkład własny"),
   ("Zestaw do mycia okien Unger (ściągaczki, kij teleskopowy, wiadro)", "sprzęt", "nowy", 1, 700, 0.23, ceneo("unger zestaw do mycia okien"), "wkład własny"),
   ("Drabina aluminiowa 3 × 9 szczebli", "sprzęt", "nowy", 1, 600, 0.23, ceneo("drabina aluminiowa 3x9"), "wkład własny"),
   ("Ulotki i wizytówki", "reklama", "nowy", 1, 450, 0.23, "[oferta drukarni – do dołączenia]", "wkład własny"),
   ("Odzież robocza i obuwie ochronne", "inne", "nowy", 1, 300, 0.23, ceneo("odziez robocza damska"), "wkład własny"),
   ("Środki czystości – druga część zapasu", "towar", "nowy", 1, 1000, 0.23, "[oferta hurtowni – do dołączenia]", "wkład własny"),
  ],
  limits=dict(sprzęt=1, towar=0.30, reklama=1, remont=0, inne=1),
  services=[
   ("Sprzątanie biura / gabinetu – wizyta ok. 2 h", 100, 24, 2.0),
   ("Sprzątanie klatki schodowej – wizyta", 90, 18, 1.5),
   ("Mycie okna dwuskrzydłowego", 45, 30, 0.4),
   ("Sprzątanie jednorazowe: mieszkanie lub po remoncie (średnio)", 450, 3, 6.0),
   ("Pranie wykładziny lub tapicerki (średnio na zlecenie)", 250, 3, 2.5),
  ],
  season=[0.85, 0.85, 1.10, 1.20, 1.05, 1.00, 0.85, 0.85, 1.00, 1.05, 1.05, 1.15],
  rozruch=[0.5, 0.7, 0.85] + [1]*9,
  koszty=[("Czynsz / najem lokalu (nie dotyczy – sprzęt w garażu własnym)", 0), ("Media (nie dotyczy)", 0),
          ("Telefon i internet", 60), ("Księgowość (biuro rachunkowe, ryczałt)", 180), ("Ubezpieczenie OC działalności", 60),
          ("Paliwo i eksploatacja auta", 400), ("Reklama bieżąca", 100),
          ("Inne (worki, rękawice, konto firmowe, drobne naprawy)", 120)],
  materialy=0.07,
 ),
}

def brutto(qty, netto, vat):
    g = round(qty * netto, 2)
    return round(g + round(g * vat, 2), 2)

def solve_adjust(cfg):
    items = [list(x) for x in cfg["items"]]
    target = cfg["target_dotacja"]
    idx = next(i for i, it in enumerate(items) if it[1] == cfg["adjust_row"] and it[7] == "dotacja")
    others = sum(brutto(it[3], it[4], it[5]) for i, it in enumerate(items) if it[7] == "dotacja" and i != idx)
    need = round(target - others, 2)
    base = round(need / (1 + items[idx][5]), 2)
    for d in [0, 0.01, -0.01, 0.02, -0.02, 0.03, -0.03]:
        cand = round(base + d, 2)
        if abs(brutto(1, cand, items[idx][5]) - need) < 0.001:
            items[idx][4] = cand
            return items
    raise SystemExit(f"nie da się trafić w {target}: need {need}")

def build(name, cfg):
    outdir = f"{REPO}/przyklady/{name}"
    tmpdir = f"{SCR}/tmp_{name}"
    os.makedirs(tmpdir, exist_ok=True)
    wb = openpyxl.load_workbook(TPL)
    z = wb["Zalozenia"]
    for k, v in cfg["zal"].items():
        z[k] = v
    k = wb["Kosztorys"]
    for r in range(5, 25):
        for col in "BCDEFHLM":
            k[f"{col}{r}"] = None
    items = solve_adjust(cfg)
    cfg["items_final"] = items
    for i, it in enumerate(items):
        r = 5 + i
        k[f"B{r}"], k[f"C{r}"], k[f"D{r}"], k[f"E{r}"], k[f"F{r}"], k[f"H{r}"], k[f"L{r}"], k[f"M{r}"] = it
    for row, cat in zip(range(33, 38), ["sprzęt", "towar", "reklama", "remont", "inne"]):
        k[f"B{row}"] = cfg["limits"][cat]
    p = wb["Przychody"]
    for i in range(5):
        r = 6 + i
        if i < len(cfg["services"]):
            n, c, q, t = cfg["services"][i]
        else:
            n, c, q, t = None, 0, 0, 0
        p[f"A{r}"], p[f"B{r}"], p[f"C{r}"], p[f"D{r}"] = n, c, q, t
    for j, col in enumerate("CDEFGHIJKLMN"):
        p[f"{col}14"] = cfg["season"][j]
        p[f"{col}19"] = cfg["rozruch"][j]
    ko = wb["Koszty"]
    dw = wb["Do-wniosku"]
    for i, (lab, val) in enumerate(cfg["koszty"]):
        ko[f"A{9+i}"] = lab
        ko[f"B{9+i}"] = val
        dw[f"A{10+i}"] = "– " + lab[0].lower() + lab[1:]
    ko["B20"] = cfg["materialy"]
    src = f"{tmpdir}/model.xlsx"
    wb.save(src)
    # przeliczenie przez LibreOffice
    conv = f"{tmpdir}/conv"
    os.makedirs(conv, exist_ok=True)
    subprocess.run(["/usr/bin/soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", conv, src],
                   check=True, capture_output=True, timeout=180)
    shutil.copy(f"{conv}/model.xlsx", f"{outdir}/model.xlsx")
    return read_values(f"{outdir}/model.xlsx")

def read_values(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = {"ark": {}, "kontrole": {}, "mies": {}, "kwart": {}, "kosztorys": [], "kategorie": {}, "przychody": {}}
    for n, d in wb.defined_names.items():
        if n.startswith("ark_"):
            sh, ref = list(d.destinations)[0]
            out["ark"][n] = wb[sh][ref.replace("$", "")].value
    z = wb["Zalozenia"]
    for r in range(27, 35):
        out["kontrole"][z[f"A{r}"].value] = z[f"B{r}"].value
    dw = wb["Do-wniosku"]
    for r in range(7, 23):
        out["mies"][dw[f"A{r}"].value] = [dw.cell(r, c).value for c in range(3, 16)]
    out["mies"]["_miesiace"] = [str(dw.cell(6, c).value)[:10] for c in range(3, 15)]
    for r in range(27, 43):
        out["kwart"][dw[f"A{r}"].value] = [dw.cell(r, c).value for c in range(3, 8)]
    k = wb["Kosztorys"]
    for r in range(5, 25):
        if k[f"B{r}"].value:
            out["kosztorys"].append({c: k[f"{c}{r}"].value for c in "ABCDEFGHIJKLM"})
    for r in range(33, 38):
        out["kategorie"][k[f"A{r}"].value] = [k[f"B{r}"].value, k[f"C{r}"].value, k[f"D{r}"].value, k[f"E{r}"].value]
    out["kosztorys_sumy"] = {"razem": k["K25"].value, "dotacja": k["K26"].value, "wklad": k["K27"].value,
                             "status": k["L29"].value, "kontrola_kat": k["C39"].value}
    p = wb["Przychody"]
    out["przychody"]["obłożenie"] = [p.cell(38, c).value for c in range(3, 15)]
    out["przychody"]["sprzedaz_szt"] = {p[f"A{r}"].value: [p.cell(r, c).value for c in range(3, 16)] for r in range(22, 27)}
    pr = wb["Prognoza"]
    out["prognoza"] = {pr[f"A{r}"].value: pr[f"C{r}"].value for r in range(24, 35)}
    out["prognoza_tab"] = [[pr.cell(r, c).value for c in range(1, 19)] for r in range(4, 20)]
    ko = wb["Koszty"]
    out["koszty_stale"] = {ko[f"A{r}"].value: ko[f"B{r}"].value for r in range(9, 17)}
    out["zus_mies"] = [ko.cell(30, c).value for c in range(3, 15)]
    out["zdrow_mies"] = [ko.cell(40, c).value for c in range(3, 15)]
    return out

if __name__ == "__main__":
    allv = {}
    for name, cfg in EXAMPLES.items():
        if len(sys.argv) > 1 and sys.argv[1] not in name:
            continue
        v = build(name, cfg)
        allv[name] = v
        a = v["ark"]
        print(name, "| przychód", round(a["ark_przychod_12m"]), "| zysk netto", round(a["ark_zysk_netto_12m"]),
              "| mies", round(a["ark_zysk_netto_mies"]), "| próg", round(a["ark_prog_rent_zl"]), a["ark_prog_rent_szt"],
              "| pes zysk", round(a["ark_zysk_netto_12m_pes"]), "| min saldo", round(a["ark_min_saldo"]), round(a["ark_min_saldo_pes"]),
              "| obł", round(a["ark_oblozenie_max"], 3), "| dotacja", v["kosztorys_sumy"])
        print("   kontrole:", v["kontrole"])
