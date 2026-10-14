#!/usr/bin/env python3
"""Generuje kosztorys-i-prognoza.md z przeliczonego model.xlsx przykładu.

Użycie:
  python3 przyklady/narzedzia/kosztorys_md.py przyklady/01-fryzjerka-mobilna-turek

Wszystkie liczby w pliku wynikowym pochodzą z arkusza (data_only=True, wartości zapisane
przez LibreOffice przy przeliczeniu). Komentarze słowne bierze z pliku narzedzia/notatki/<katalog-przykładu>.md
(bloki zaczynające się od '@@ nazwa').
"""
import os, re, sys
import openpyxl

def zl(x, dec=0):
    if x is None:
        return "–"
    s = f"{round(x + (1e-9 if x >= 0 else -1e-9), dec):,.{dec}f}".replace(",", " ").replace(".", ",")
    return s

def pct(x, dec=0):
    return f"{round(x * 100, dec):.{dec}f}".replace(".", ",") + "%"

def read(folder):
    wb = openpyxl.load_workbook(os.path.join(folder, "model.xlsx"), data_only=True)
    return wb

def notes(folder):
    p = os.path.join(os.path.dirname(os.path.abspath(__file__)), "notatki", os.path.basename(os.path.normpath(folder)) + ".md")
    out = {}
    if os.path.exists(p):
        cur = None
        for line in open(p, encoding="utf-8"):
            m = re.match(r"^@@ (\w+)", line)
            if m:
                cur = m.group(1); out[cur] = []
                continue
            if cur:
                out[cur].append(line.rstrip("\n"))
    return {k: "\n".join(v).strip() for k, v in out.items()}

def main(folder):
    wb = read(folder)
    N = notes(folder)
    z, k, p, ko, pr, dw = (wb[s] for s in ["Zalozenia", "Kosztorys", "Przychody", "Koszty", "Prognoza", "Do-wniosku"])
    L = []
    a = L.append
    a(N.get("naglowek", "# Kosztorys i prognoza"))
    a("")
    a("## 1. Założenia wpisane do arkusza (`model.xlsx`, arkusz `Zalozenia`)")
    a("")
    a("| Parametr | Wartość |")
    a("|---|---|")
    for r in range(4, 23):
        v = z[f"B{r}"].value
        if hasattr(v, "strftime"):
            v = v.strftime("%d.%m.%Y")
        elif isinstance(v, float) and v < 1:
            v = pct(v, 1).replace(",0%", "%")
        elif isinstance(v, (int, float)):
            v = zl(v, 2) if r in (5, 6, 7, 16) else zl(v)
        a(f"| {z[f'A{r}'].value} | {v} |")
    a("")
    a("## 2. Specyfikacja wydatków (arkusz `Kosztorys`)")
    a("")
    a(N.get("kosztorys_intro", ""))
    a("")
    a("| Lp. | Pozycja | Kategoria | Nowy/używany | Ilość | Cena jedn. netto | Wartość netto | VAT | Brutto | Do rozliczenia | Finansowanie | Oferta typu |")
    a("|---|---|---|---|---|---|---|---|---|---|---|---|")
    for r in range(5, 25):
        if not k[f"B{r}"].value:
            continue
        link = k[f"L{r}"].value or ""
        if link.startswith("http"):
            host = re.sub(r"^https?://(www\.)?", "", link).split("/")[0]
            link = f"[{host}]({link})"
        a(f"| {k[f'A{r}'].value} | {k[f'B{r}'].value} | {k[f'C{r}'].value} | {k[f'D{r}'].value} | {k[f'E{r}'].value} | "
          f"{zl(k[f'F{r}'].value, 2)} | {zl(k[f'G{r}'].value, 2)} | {zl(k[f'I{r}'].value, 2)} | {zl(k[f'J{r}'].value, 2)} | "
          f"{zl(k[f'K{r}'].value, 2)} | {k[f'M{r}'].value} | {link} |")
    a(f"| | **RAZEM** | | | | | {zl(k['G25'].value, 2)} | {zl(k['I25'].value, 2)} | {zl(k['J25'].value, 2)} | **{zl(k['K25'].value, 2)}** | | |")
    a(f"| | w tym z dotacji | | | | | | | | **{zl(k['K26'].value, 2)}** | dotacja | |")
    a(f"| | w tym z wkładu własnego | | | | | | | | **{zl(k['K27'].value, 2)}** | wkład własny | |")
    a("")
    a(f"Kontrola arkusza: {k['L29'].value}; suma kategorii = suma z dotacji: {k['C39'].value}.")
    a("")
    a("### Limity kategorii według regulaminu urzędu")
    a("")
    a("| Kategoria | Limit wpisany do arkusza | Suma z dotacji (zł) | Udział w dotacji | Status |")
    a("|---|---|---|---|---|")
    for r in range(33, 38):
        lim = k[f"B{r}"].value
        lim_s = "brak limitu %" if lim >= 1 else ("wykluczone" if lim == 0 else pct(lim))
        a(f"| {k[f'A{r}'].value} | {lim_s} | {zl(k[f'C{r}'].value, 2)} | {pct(k[f'D{r}'].value, 1)} | {k[f'E{r}'].value} |")
    a("")
    a(N.get("limity", ""))
    a("")
    a("## 3. Przychody – założenia (arkusz `Przychody`)")
    a("")
    a("| Usługa | Cena (zł) | Liczba / mies. przy pełnym obłożeniu | Czas 1 usługi (h) |")
    a("|---|---|---|---|")
    for r in range(6, 11):
        if p[f"A{r}"].value:
            a(f"| {p[f'A{r}'].value} | {zl(p[f'B{r}'].value)} | {p[f'C{r}'].value} | {str(p[f'D{r}'].value).replace('.', ',')} |")
    a("")
    months = [dw.cell(6, c).value for c in range(3, 15)]
    mlab = [m.strftime("%m.%Y") if hasattr(m, "strftime") else str(m) for m in months]
    a("| | " + " | ".join(mlab) + " |")
    a("|---|" + "---|" * 12)
    a("| Rozruch | " + " | ".join(pct(p.cell(19, c).value) for c in range(3, 15)) + " |")
    a("| Sezonowość | " + " | ".join(str(p.cell(20, c).value).replace(".", ",") for c in range(3, 15)) + " |")
    for r in range(22, 27):
        if p[f"A{r}"].value:
            a(f"| {p[f'A{r}'].value} (szt.) | " + " | ".join(str(p.cell(r, c).value) for c in range(3, 15)) + " |")
    a("| Obłożenie czasu pracy | " + " | ".join(pct(p.cell(38, c).value) for c in range(3, 15)) + " |")
    a("")
    a(N.get("przychody", ""))
    a("")
    a("## 4. Koszty stałe miesięcznie (arkusz `Koszty`)")
    a("")
    a("| Pozycja | zł / mies. |")
    a("|---|---|")
    for r in range(9, 17):
        a(f"| {ko[f'A{r}'].value} | {zl(ko[f'B{r}'].value)} |")
    a(f"| **Razem koszty stałe** | **{zl(ko['C17'].value)}** |")
    a(f"| Materiały (koszt zmienny) | {pct(ko['B20'].value)} przychodu |")
    a("")
    a(N.get("koszty", ""))
    a("")
    a("## 5. Kalkulacja 12 miesięcy – układ kwartalny (arkusz `Do-wniosku`, tabela 1a)")
    a("")
    a("| Pozycja | Kw. I (mies. 1–3) | Kw. II (mies. 4–6) | Kw. III (mies. 7–9) | Kw. IV (mies. 10–12) | Razem 12 mies. |")
    a("|---|---|---|---|---|---|")
    for r in range(27, 43):
        lab = dw[f"A{r}"].value
        vals = [dw.cell(r, c).value for c in range(3, 8)]
        if all((v or 0) == 0 for v in vals):
            continue
        b = r in (27, 28, 40, 42)
        cells = [zl(v) for v in vals]
        if b:
            lab = f"**{lab}**"; cells = [f"**{c}**" for c in cells]
        a(f"| {lab} | " + " | ".join(cells) + " |")
    a(f"\nKontrola arkusza: suma kwartałów = suma miesięcy – {dw['G43'].value}.")
    a("")
    a("### Ten sam rachunek miesiąc po miesiącu (wymóg rozporządzenia: przychody i koszty w 1. roku miesięcznie)")
    a("")
    a("| Pozycja | " + " | ".join(mlab) + " | Razem |")
    a("|---|" + "---|" * 13)
    for r in (7, 8, 18, 19, 20, 21, 22):
        a(f"| {dw[f'A{r}'].value} | " + " | ".join(zl(dw.cell(r, c).value) for c in range(3, 16)) + " |")
    a("")
    a("## 6. Wskaźniki i próg rentowności (arkusze `Prognoza` i `Do-wniosku`, tabela 3)")
    a("")
    a("| Wskaźnik | Scenariusz bazowy | Scenariusz −30% przychodów |")
    a("|---|---|---|")
    for r in range(75, 88):
        lab = dw[f"A{r}"].value
        c, d = dw[f"C{r}"].value, dw[f"D{r}"].value
        if isinstance(c, float) and abs(c) < 5 and "Rentowność" in lab or "Zwrot z dotacji" in lab:
            cs, ds = pct(c, 1), pct(d, 1)
        elif "Okres zwrotu" in lab:
            cs = zl(c, 1) if isinstance(c, (int, float)) else c
            ds = zl(d, 1) if isinstance(d, (int, float)) else d
        elif "liczba sprzedaży" in lab:
            cs, ds = str(c), str(d)
        else:
            cs, ds = zl(c), zl(d)
        a(f"| {lab} | {cs} | {ds} |")
    a(f"| Próg rentowności po okresie ulg ZUS (pełny ZUS) – przychód mies. (zł) | {zl(pr['C29'].value)} | – |")
    a(f"| Średni przychód miesięczny od 4. miesiąca (zł) | {zl(pr['C30'].value)} | {zl(pr['C32'].value)} |")
    a(f"| Margines bezpieczeństwa (średni przychód od m4 vs próg) | {pct(pr['C31'].value)} | {pct(pr['C33'].value)} |")
    a(f"| Pierwszy miesiąc z dodatnim wynikiem netto | {pr['C34'].value} | – |")
    a(f"| Najwyższe obłożenie czasu pracy | {pct(z['B33'].value)} | – |")
    a("")
    a("Próg rentowności = średnie miesięczne koszty stałe z ZUS i składką zdrowotną "
      f"({zl(pr['C24'].value)} zł) ÷ (1 − udział materiałów {pct(pr['C25'].value)}). "
      f"Średnia cena jednej usługi w roku: {zl(pr['C27'].value, 2)} zł, stąd liczba usług w progu.")
    a("")
    a("## 7. Scenariusz pesymistyczny (przychody o 30% niższe w każdym miesiącu)")
    a("")
    a(N.get("pesymistyczny", ""))
    a("")
    a("## 8. Kontrole spójności arkusza (`Zalozenia`, wiersze 27–34)")
    a("")
    a("| Kontrola | Wynik |")
    a("|---|---|")
    for r in range(27, 35):
        v = z[f"B{r}"].value
        if isinstance(v, float):
            v = pct(v) if r == 33 else zl(v, 2)
        a(f"| {z[f'A{r}'].value} | {v} |")
    a("")
    a(N.get("uproszczenia", ""))
    a("")
    out = os.path.join(folder, "kosztorys-i-prognoza.md")
    open(out, "w", encoding="utf-8").write("\n".join(L).replace("\n\n\n", "\n\n"))
    print("zapisano", out)

if __name__ == "__main__":
    main(sys.argv[1])
