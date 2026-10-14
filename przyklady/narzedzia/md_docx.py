#!/usr/bin/env python3
"""Prosty konwerter Markdown → DOCX (python-docx) dla materiałów klienta.

Użycie:
  python3 przyklady/narzedzia/md_docx.py WEJŚCIE.md WYJŚCIE.docx

Obsługuje: nagłówki #/##/###, akapity, listy (- / –), listy numerowane (numer zostaje w tekście),
tabele, cytaty (>), bloki ```…``` (np. wzory pism – wstawiane w ramce do edycji), **pogrubienia**.
"""
import re, sys
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Cm

def shade(obj, fill):
    el = obj._tc if hasattr(obj, "_tc") else obj._p
    pr = el.get_or_add_tcPr() if hasattr(obj, "_tc") else el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pr.append(shd)

def runs(par, text, size=None, italic=False):
    for part in re.split(r"(\*\*.+?\*\*)", text):
        if not part:
            continue
        b = part.startswith("**") and part.endswith("**")
        r = par.add_run((part[2:-2] if b else part).replace("`", ""))
        r.bold = b; r.italic = italic
        if size:
            r.font.size = Pt(size)

def table(doc, rows):
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [r for r in cells if not all(re.fullmatch(r":?-{2,}:?", c) for c in r)]
    n = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=n); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(cells):
        row = t.add_row()
        for j in range(n):
            c = row.cells[j]; c.text = ""
            runs(c.paragraphs[0], r[j] if j < len(r) else "", size=9)
            if i == 0:
                for x in c.paragraphs[0].runs:
                    x.bold = True
                shade(c, "D9D9D9")
        if i == 0:
            tr = row._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); tr.append(h)
    doc.add_paragraph()

def frame(doc, lines):
    t = doc.add_table(rows=1, cols=1); t.style = "Table Grid"
    c = t.rows[0].cells[0]; c.text = ""
    first = True
    for line in lines:
        p = c.paragraphs[0] if first else c.add_paragraph()
        first = False
        runs(p, line if line.strip() else " ", size=10)
    doc.add_paragraph()

def convert(src, dst):
    md = open(src, encoding="utf-8").read().splitlines()
    doc = Document()
    s = doc.sections[0]
    s.page_height, s.page_width = Cm(29.7), Cm(21.0)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(2))
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(11)
    tb, code, in_code = [], [], False
    for line in md:
        if line.startswith("```"):
            if in_code:
                frame(doc, code); code = []; in_code = False
            else:
                if tb:
                    table(doc, tb); tb = []
                in_code = True
            continue
        if in_code:
            code.append(line); continue
        if line.startswith("|"):
            tb.append(line); continue
        if tb:
            table(doc, tb); tb = []
        if line.startswith("# "):
            doc.add_heading(line[2:], 0)
        elif line.startswith("## "):
            doc.add_heading(line[3:], 1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], 2)
        elif line.startswith("> "):
            p = doc.add_paragraph(); runs(p, line[2:], size=10, italic=True); shade(p, "F2F2F2")
        elif line.startswith("- ") or line.startswith("– "):
            p = doc.add_paragraph(style="List Bullet"); runs(p, line[2:])
        elif line.strip() in ("---", ""):
            continue
        else:
            p = doc.add_paragraph(); runs(p, line)
    if tb:
        table(doc, tb)
    doc.save(dst)
    print("zapisano", dst)

if __name__ == "__main__":
    convert(sys.argv[1], sys.argv[2])
