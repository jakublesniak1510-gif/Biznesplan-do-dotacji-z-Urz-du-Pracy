"""Pełny przebieg `pup-gen generuj --dry-run` na trzech przykładach (prawdziwy LibreOffice)."""

import json

import pytest
from docx import Document

from pup_gen import KATALOG_PRZYKLADOW
from pup_gen.cli import main
from pup_gen.schemat import WynikGeneracji

from conftest import PRZYKLADY, mapa

pytestmark = pytest.mark.soffice


def _generuj(nazwa, plik_mapy, out, archiwum, *dodatkowe):
    kod = main(["generuj", "--kwestionariusz", str(KATALOG_PRZYKLADOW / f"kwestionariusz-{nazwa}.json"),
                "--mapa", str(KATALOG_PRZYKLADOW / f"{plik_mapy}.json"), "--dry-run",
                "--out", str(out), "--archiwum", str(archiwum), *dodatkowe])
    assert kod == 0
    return WynikGeneracji.model_validate_json((out / "wynik.json").read_text(encoding="utf-8"))


@pytest.mark.parametrize("nazwa", PRZYKLADY)
def test_dry_run_przyklad_dluga_mapa(nazwa, tmp_path):
    out = tmp_path / "wyniki" / nazwa
    wynik = _generuj(nazwa, "mapa-dluga-14-pol", out, tmp_path / "puste-archiwum")
    m = mapa("mapa-dluga-14-pol")
    docx = next(out.glob("PUP-*_wniosek_v1.docx"))
    doc = Document(str(docx))
    naglowki = [p.text for p in doc.paragraphs if p.style.name.startswith("Heading")]
    for pole in m.pola:  # każde pole mapy, z nazwą jak w formularzu, w kolejności formularza
        assert f"{pole.id}. {pole.nazwa_w_formularzu}" in naglowki
    kolejnosc = [naglowki.index(f"{p.id}. {p.nazwa_w_formularzu}") for p in m.pola]
    assert kolejnosc == sorted(kolejnosc)
    assert len(doc.tables) == 3  # kosztorys, kalkulacja miesięczna, wskaźniki
    tekst = "\n".join(p.text for p in doc.paragraphs)
    assert "Co wpisujesz i podpisujesz sam(a)" in tekst
    assert "Raport kontroli" not in tekst and "BLOKUJE" not in tekst  # raport tylko w osobnym pliku
    assert "{{" not in tekst
    for p in wynik.pola:
        if p.limit_znakow and p.tekst:
            assert len(p.tekst) <= p.limit_znakow
    raport = json.loads((out / "raport-kontroli.json").read_text(encoding="utf-8"))
    assert raport["blokuje"] == 0, (out / "raport-kontroli-WLASCICIEL.md").read_text(encoding="utf-8")
    assert (out / "raport-kontroli-WLASCICIEL.docx").exists()
    assert (out / f"arkusz-{wynik.id_zlecenia}.xlsx").exists()
    assert main(["sprawdz", "--wynik", str(out), "--archiwum", str(tmp_path / "puste-archiwum")]) == 0


def test_krotka_mapa_i_pdf(tmp_path):
    out = tmp_path / "krotka"
    wynik = _generuj("firma-sprzatajaca", "mapa-krotka-6-pol", out, tmp_path / "brak", "--pdf")
    assert [p.id for p in wynik.pola] == ["1", "2", "3", "4", "5", "6"]
    assert all(p.tekst for p in wynik.pola)
    assert next(out.glob("*.pdf")).stat().st_size > 10_000
    doc = Document(str(next(out.glob("*.docx"))))
    assert len(doc.tables) == 2  # kosztorys + kalkulacja kwartalna


def test_drugi_plan_do_tego_samego_pup_ostrzega_o_podobienstwie(tmp_path):
    archiwum = tmp_path / "wyniki"
    _generuj("fryzjerka-mobilna", "mapa-krotka-6-pol", archiwum / "zielonow" / "pierwszy", archiwum)
    # ten sam klient pod innym id = „bliźniak”
    dane = json.loads((KATALOG_PRZYKLADOW / "kwestionariusz-fryzjerka-mobilna.json").read_text(encoding="utf-8"))
    dane["id"] = "2026-10-blizniak"
    plik = tmp_path / "blizniak.json"
    plik.write_text(json.dumps(dane, ensure_ascii=False), encoding="utf-8")
    out = archiwum / "zielonow" / "drugi"
    assert main(["generuj", "--kwestionariusz", str(plik), "--mapa", str(KATALOG_PRZYKLADOW / "mapa-krotka-6-pol.json"),
                 "--dry-run", "--out", str(out), "--archiwum", str(archiwum)]) == 0
    raport = json.loads((out / "raport-kontroli.json").read_text(encoding="utf-8"))
    podob = [u for u in raport["ustalenia"] if u["kategoria"] == "podobieństwo"]
    assert podob and podob[0]["poziom"] == "OSTRZEŻENIE" and "2026-10-fryzjerka-anna" in podob[0]["komunikat"]


def test_bledny_kwestionariusz_konczy_sie_komunikatem(tmp_path, capsys):
    dane = json.loads((KATALOG_PRZYKLADOW / "kwestionariusz-fryzjerka-mobilna.json").read_text(encoding="utf-8"))
    dane["notatka_wlasciciela"] = "PESEL 44051401359"
    plik = tmp_path / "zly.json"
    plik.write_text(json.dumps(dane, ensure_ascii=False), encoding="utf-8")
    with pytest.raises(SystemExit) as e:
        main(["generuj", "--kwestionariusz", str(plik), "--mapa", str(KATALOG_PRZYKLADOW / "mapa-krotka-6-pol.json"),
              "--dry-run", "--out", str(tmp_path / "x")])
    assert "PESEL" in str(e.value.code) and "44051401359" not in str(e.value.code)
