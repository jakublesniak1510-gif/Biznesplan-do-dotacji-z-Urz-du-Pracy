import pytest
from docx import Document

from pup_gen import finanse
from pup_gen.mapowanie import mapuj
from pup_gen.schemat import MapaFormularza


def _formularz(plik):
    d = Document()
    d.add_paragraph("WNIOSEK O PRZYZNANIE ŚRODKÓW NA PODJĘCIE DZIAŁALNOŚCI GOSPODARCZEJ")
    d.add_paragraph("I. Dane wnioskodawcy")
    d.add_paragraph("1. Opis planowanego przedsięwzięcia (max. 2000 znaków)")
    d.add_paragraph("2. Kwalifikacje i doświadczenie zawodowe")
    d.add_paragraph("Proszę opisać wykształcenie i kursy – maksymalnie 1 strona.")
    d.add_paragraph("3. Analiza rynku i konkurencji (do 3 000 znaków)")
    d.add_paragraph("4. Sposób pozyskania klientów i promocji – do 300 słów")
    d.add_paragraph("5. Co Pan/i zrobi, jeśli przychody będą niższe od planowanych?")
    d.add_paragraph("6. Proponowana forma zabezpieczenia zwrotu")
    d.save(plik)
    return plik


def test_mapuj_docx(tmp_path):
    wynik = mapuj(_formularz(tmp_path / "wniosek.docx"), pup="PUP testowy")
    pola = {p["id"]: p for p in wynik["pola"]}
    assert pola["1"]["sekcje"] == ["01"] and pola["1"]["limit_znakow"] == 1900
    assert pola["2"]["sekcje"] == ["02"] and pola["2"]["limit_znakow"] == 1620
    assert pola["3"]["sekcje"] == ["03", "04"]
    assert pola["3"]["limit_znakow"] == 2850 and pola["3"]["podzial_limitu"]
    assert pola["4"]["sekcje"][0] == "05" and pola["4"]["limit_znakow"] == 1890
    assert pola["5"]["sekcje"] == ["09"] and pola["5"]["uwagi"].endswith("?")
    assert any("zabezpieczenia" in x for x in wynik["poza_zakresem"])
    assert any("Dane wnioskodawcy" in x for x in wynik["poza_zakresem"])
    assert any("Pole 5" in x for x in wynik["do_wyjasnienia"])
    MapaFormularza.model_validate(wynik)  # wynik nadaje się do generowania po korekcie


@pytest.mark.soffice
def test_mapuj_pdf_plaski(tmp_path):
    pdf = finanse.konwertuj_pdf(_formularz(tmp_path / "wniosek.docx"), tmp_path)
    wynik = mapuj(pdf)
    assert wynik["typ_pliku"] == "pdf_plaski"
    assert {p["id"] for p in wynik["pola"]} >= {"1", "2", "3", "4", "5"}


def test_mapuj_nieobslugiwany_format(tmp_path):
    plik = tmp_path / "a.odt"
    plik.write_text("x")
    with pytest.raises(ValueError):
        mapuj(plik)
