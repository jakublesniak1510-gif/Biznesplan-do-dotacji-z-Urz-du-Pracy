"""Heurystyczne mapowanie formularza PUP (PDF/DOCX) na pola z limitami – wynik do ręcznej korekty.

Reguły (słowa kluczowe, przeliczanie limitów, współczynniki 0,95 / 0,90) według
`szablony/10-mapowanie-pol-formularza.md`. Pliki od klienta są niezaufane: tylko je czytamy.
"""

from __future__ import annotations

import re
import shutil
import subprocess
from pathlib import Path
from typing import Any

SLOWNIK_SEKCJI: list[tuple[str, tuple[str, ...]]] = [
    ("09", ("ryzyk", "zagrożenia dla realizacji", "plan awaryjny", "co w przypadku", "jeśli przychody", "niższe od planowanych")),
    ("07", ("swot", "mocne i słabe", "mocne strony", "szanse")),
    ("04", ("konkuren", "przewag")),
    ("03", ("rynek", "rynku", "klienci", "klientów docelow", "odbiorc", "popyt", "zapotrzebowan", "zasięg", "sezonow")),
    ("05", ("marketing", "promocj", "reklam", "pozyskani", "pozyskania klientów", "sprzedaży", "dystrybuc")),
    ("06", ("specyfikacj", "kosztorys", "zakup", "uzasadnienie wydatk", "wydatków", "sprzęt", "wyposażeni")),
    ("08", ("przychod", "kalkulacj", "efektywnoś", "zysk", "opłacalnoś", "próg rentown", "prognoz", "założeń finansow")),
    ("02", ("kwalifikac", "wykształceni", "doświadczeni", "przygotowanie zawodowe", "kursy", "uprawnieni")),
    ("01", ("opis przedsięwzięcia", "opis planowan", "charakterystyka działalności", "przedmiot", "rodzaj działalności",
            "pkd", "lokalizacj", "miejsce prowadzenia", "dlaczego", "motywac")),
]
POZA_ZAKRESEM = ("dane osobowe", "dane wnioskodawcy", "pesel", "adres", "rachunek bankow", "zabezpieczeni",
                 "poręcz", "weksel", "oświadcz", "de minimis", "załącznik", "podpis")
_RE_LIMIT = re.compile(
    r"(max\.?|maks\.?|maksymalnie|do|nie więcej niż|limit)\s*(\d[\d\s]*)\s*(znak\w*|słów|słowa|wyraz\w*|stron\w*|lini\w*|wiersz\w*)"
    r"(\s*bez\s+spacji)?", re.I)
_RE_POL_STRONY = re.compile(r"pół\s+strony", re.I)
_RE_NAGLOWEK = re.compile(r"^\s*((?:[IVX]+\.)?\d+(?:\.\d+)*\.?|[IVX]+\.|[a-h]\))\s+(.{5,220})$")


def _limit(tekst: str) -> tuple[dict[str, Any] | None, int | None]:
    """Zwraca (surowy zapis limitu, limit_znakow do generowania)."""
    if _RE_POL_STRONY.search(tekst):
        return {"wartosc": 0.5, "jednostka": "strony", "zrodlo": "„pół strony”"}, 810
    m = _RE_LIMIT.search(tekst)
    if not m:
        return None, None
    wartosc = int(re.sub(r"\s", "", m.group(2)))
    jedn = m.group(3).lower()
    if jedn.startswith("znak"):
        if m.group(4):
            jednostka, znaki, wsp = "znaki_bez_spacji", wartosc * 1.15, 0.90
        else:
            jednostka, znaki, wsp = "znaki_ze_spacjami", wartosc, 0.95
    elif jedn.startswith(("słow", "słów", "wyraz")):
        jednostka, znaki, wsp = "slowa", wartosc * 7, 0.90
    elif jedn.startswith("stron"):
        jednostka, znaki, wsp = "strony", wartosc * 1800, 0.90
    else:
        jednostka, znaki, wsp = "linie", wartosc * 85, 0.90
    return ({"wartosc": wartosc, "jednostka": jednostka, "zrodlo": f"„{m.group(0).strip()}”"},
            int(znaki * wsp) // 10 * 10)


def _zawiera(tekst: str, slowa: tuple[str, ...]) -> bool:
    """Słowo kluczowe na początku wyrazu („zysk” nie pasuje do „pozyskania”, „oświadcz” do „doświadczenia”)."""
    t = tekst.lower()
    return any(re.search(r"(?<!\w)" + re.escape(w), t) for w in slowa)


def _sekcje(tekst: str) -> list[str]:
    """Pierwsza pasująca sekcja; dwie tylko dla pola łączonego „rynek + konkurencja” (03 + 04)."""
    trafione = [s for s, slowa in SLOWNIK_SEKCJI if _zawiera(tekst, slowa)]
    if {"03", "04"} <= set(trafione):
        return ["03", "04"]
    return trafione[:1]


def _linie_docx(plik: Path) -> list[tuple[int, str]]:
    from docx import Document

    doc = Document(str(plik))
    linie = [(1, p.text.strip()) for p in doc.paragraphs if p.text.strip()]
    for t in doc.tables:
        for wiersz in t.rows:
            for c in wiersz.cells:
                for akapit in c.text.splitlines():
                    if akapit.strip():
                        linie.append((1, akapit.strip()))
    return linie


def _pola_acroform(plik: Path) -> list[dict[str, Any]]:
    try:
        from pypdf import PdfReader
    except ImportError:
        return []
    pola = []
    for nazwa, f in (PdfReader(str(plik)).get_fields() or {}).items():
        if f.get("/FT") == "/Tx":
            pola.append({"nazwa": nazwa, "max_len": f.get("/MaxLen"), "opis": f.get("/TU")})
    return pola


def _linie_pdf(plik: Path) -> list[tuple[int, str]]:
    teksty: list[str] = []
    try:
        from pypdf import PdfReader
        teksty = [s.extract_text() or "" for s in PdfReader(str(plik)).pages]
    except ImportError:
        if shutil.which("pdftotext"):
            wynik = subprocess.run(["pdftotext", "-layout", str(plik), "-"], capture_output=True, text=True, timeout=60)
            teksty = wynik.stdout.split("\f")
    return [(nr, l.strip()) for nr, strona in enumerate(teksty, 1) for l in strona.splitlines() if l.strip()]


def mapuj(plik: Path, pup: str = "", formularz: str = "") -> dict[str, Any]:
    """Zwraca słownik zgodny z MapaFormularza (+ pola pomocnicze) – do korekty przez właściciela."""
    rozszerzenie = plik.suffix.lower()
    acro: list[dict[str, Any]] = []
    if rozszerzenie == ".docx":
        linie, typ = _linie_docx(plik), "docx"
    elif rozszerzenie == ".pdf":
        acro = _pola_acroform(plik)
        linie, typ = _linie_pdf(plik), ("pdf_formularz" if acro else "pdf_plaski")
    else:
        raise ValueError("Obsługiwane formaty: .pdf, .docx (inne skonwertuj: soffice --headless --convert-to docx)")
    if typ != "docx" and not linie:
        raise ValueError("PDF bez warstwy tekstowej (skan?) – najpierw OCR: ocrmypdf -l pol plik.pdf wynik.pdf")

    pola: list[dict[str, Any]] = []
    poza: list[str] = []
    do_wyjasnienia: list[str] = []
    for i, (strona, linia) in enumerate(linie):
        m = _RE_NAGLOWEK.match(linia)
        if not m:
            continue
        id_pola, nazwa = m.group(1).rstrip("."), m.group(2).strip()
        kontekst = " ".join(l for _, l in linie[i:i + 4])
        if _zawiera(nazwa, POZA_ZAKRESEM):
            poza.append(f"{id_pola}. {nazwa}")
            continue
        sekcje = _sekcje(nazwa) or _sekcje(kontekst)
        if not sekcje:
            continue
        surowy, limit = _limit(kontekst)
        pole: dict[str, Any] = {"id": id_pola, "nazwa_w_formularzu": _RE_LIMIT.sub("", nazwa).strip(" ():–-"),
                                "strona": strona, "limit": surowy, "limit_znakow": limit, "sekcje": sekcje,
                                "podzial_limitu": {sekcje[0]: 0.55, sekcje[1]: 0.45} if len(sekcje) == 2 else None,
                                "format": "tekst", "uwagi": nazwa if nazwa.endswith("?") else ""}
        if re.search(r"kalkulacj|harmonogram", nazwa, re.I) and re.search(r"tabel|miesi|kwarta", kontekst, re.I):
            pole["sekcje"] = ["tabela:kalkulacja_kwartalna" if "kwarta" in kontekst.lower() else "tabela:kalkulacja_miesieczna"]
            pole["format"], pole["podzial_limitu"], pole["limit_znakow"] = "tabela", None, None
        if limit is None and pole["format"] == "tekst":
            pole["limit_znakow"] = 2000
            do_wyjasnienia.append(f"Pole {id_pola} bez rozpoznanego limitu – przyjęto 2000 znaków")
        if len(sekcje) == 2:
            do_wyjasnienia.append(f"Pole {id_pola} łączy sekcje {sekcje} – sprawdź podział limitu")
        pola.append(pole)

    if acro:
        do_wyjasnienia.append("Pola AcroForm (nazwa, /MaxLen) – przypisz limity: " + "; ".join(
            f"{a['nazwa']}={a['max_len']}" for a in acro if a["max_len"]))
    if not pola:
        do_wyjasnienia.append("Nie rozpoznano żadnego pola – uzupełnij mapę ręcznie (szablony/10).")
    obecne = {s for p in pola for s in p["sekcje"]}
    brakujace = [s for s, _ in SLOWNIK_SEKCJI if s not in obecne]
    if brakujace:
        do_wyjasnienia.append(f"Sekcje bez pola w formularzu: {sorted(brakujace)} – świadomie pominięte?")
    return {
        "status": "DO RĘCZNEJ KOREKTY – sprawdź limity, sekcje i kolejność przed generowaniem",
        "pup": pup, "formularz": formularz or plik.stem, "plik_zrodlowy": plik.name, "typ_pliku": typ,
        "pola": pola, "tabele_z_arkusza": [], "poza_zakresem": poza, "do_wyjasnienia": do_wyjasnienia,
        "pola_acroform": acro,
    }
