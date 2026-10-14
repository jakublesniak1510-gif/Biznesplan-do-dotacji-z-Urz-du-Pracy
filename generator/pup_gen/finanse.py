"""Arkusz finansowy: wypełnienie kopii `model-finansowy-pup.xlsx`, przeliczenie w LibreOffice,
odczyt nazwanych komórek `ark_*`, kontroli spójności (`Zalozenia!B27:B34`) i tabel `Do-wniosku`.

Wpisujemy wyłącznie komórki wejściowe (żółte) – formuły zostają nietknięte.
"""

from __future__ import annotations

import logging
import os
import shutil
import subprocess
import tempfile
from dataclasses import dataclass, field
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

import openpyxl

from . import PLIK_MODELU
from .schemat import Kwestionariusz, Ostrzezenie, RegulaminPUP

log = logging.getLogger(__name__)

KATEGORIE = ("sprzęt", "towar", "reklama", "remont", "inne")  # kolejność wierszy Kosztorys!A33:A37
WIERSZ_ZAKUPOW, MAKS_ZAKUPOW = 5, 20
WIERSZ_LIMITOW = 33
WIERSZ_USLUG, MAKS_USLUG = 6, 5
KOLUMNY_MIESIECY = [openpyxl.utils.get_column_letter(c) for c in range(3, 15)]  # C..N
KOSZTY_STALE = ("czynsz", "media", "telefon_internet", "ksiegowosc", "ubezpieczenie", "paliwo",
                "reklama", "inne")  # Koszty!B9:B16

NAZWY_PROCENTOWE = ("ark_margines", "ark_rentownosc", "ark_zwrot_z_dotacji", "ark_materialy_pct",
                    "ark_oblozenie_max")
NAZWY_LICZBOWE = ("ark_prog_rent_szt", "ark_pierwszy_mies_zysku")
NAZWY_Z_GROSZAMI = ("ark_zus_m12",)


class BladArkusza(RuntimeError):
    pass


@dataclass
class WynikFinansowy:
    plik: Path
    wartosci: dict[str, float | str | None]
    ark: dict[str, str]
    kontrole: dict[str, float | str | None]
    ostrzezenia: list[Ostrzezenie] = field(default_factory=list)
    tabele: dict[str, list[list[str]]] = field(default_factory=dict)


# --------------------------------------------------------------------------- formatowanie PL
def _zaokraglij(x: float, miejsca: int = 0) -> Decimal:
    return Decimal(str(x)).quantize(Decimal(1).scaleb(-miejsca), rounding=ROUND_HALF_UP)


def liczba_pl(x: float, miejsca: int = 0) -> str:
    """12 500,00 – spacja jako separator tysięcy, przecinek dziesiętny."""
    d = _zaokraglij(x, miejsca)
    znak = "-" if d < 0 else ""
    calk, _, ulamek = f"{abs(d):.{miejsca}f}".partition(".")
    grupy = []
    while calk:
        grupy.insert(0, calk[-3:])
        calk = calk[:-3]
    return znak + " ".join(grupy) + ("," + ulamek if ulamek else "")


def procent_pl(x: float) -> str:
    tekst = liczba_pl(x * 100, 1)
    return (tekst[:-2] if tekst.endswith(",0") else tekst) + "%"


def formatuj_ark(nazwa: str, wartosc: float | str | None) -> str:
    """Wartość zmiennej `ark_*` w zapisie do tekstu wniosku."""
    if wartosc is None:
        return "[BRAK DANYCH: wartość z arkusza]"
    if isinstance(wartosc, str):
        return wartosc
    if nazwa.startswith(NAZWY_PROCENTOWE):
        return procent_pl(wartosc)
    if nazwa.startswith(NAZWY_LICZBOWE):
        return liczba_pl(wartosc)
    if nazwa.startswith("ark_okres_zwrotu_mies"):
        return f"{liczba_pl(wartosc, 1)} mies."
    if nazwa.startswith(NAZWY_Z_GROSZAMI):
        return f"{liczba_pl(wartosc, 2)} zł"
    return f"{liczba_pl(wartosc)} zł"


def formatuj_komorke(wartosc: Any, format_liczby: str) -> str:
    if wartosc is None:
        return ""
    if isinstance(wartosc, (int, float)) and not isinstance(wartosc, bool):
        if "%" in format_liczby:
            return procent_pl(wartosc)
        if "0.00" in format_liczby:
            return liczba_pl(wartosc, 2)
        if "0.0" in format_liczby:
            return liczba_pl(wartosc, 1)
        if format_liczby.startswith("mm"):
            return str(wartosc)
        return liczba_pl(wartosc)
    if hasattr(wartosc, "strftime"):
        return wartosc.strftime("%m.%Y")
    return str(wartosc)


# --------------------------------------------------------------------------- wypełnienie
def wypelnij(kw: Kwestionariusz, regulamin: RegulaminPUP | None, plik: Path,
             szablon: Path = PLIK_MODELU) -> list[Ostrzezenie]:
    """Zapisuje do `plik` kopię szablonu z danymi klienta. Zwraca ostrzeżenia o przyjętych założeniach."""
    ostrz: list[Ostrzezenie] = []

    def uwaga(poziom: str, komunikat: str) -> None:
        ostrz.append(Ostrzezenie(poziom=poziom, kategoria="arkusz – założenia", komunikat=komunikat))

    if len(kw.zakupy) > MAKS_ZAKUPOW:
        raise BladArkusza(f"Arkusz mieści {MAKS_ZAKUPOW} pozycji kosztorysu, podano {len(kw.zakupy)}.")
    wb = openpyxl.load_workbook(szablon)
    z, ks, pr, ko = wb["Zalozenia"], wb["Kosztorys"], wb["Przychody"], wb["Koszty"]

    # --- Zalozenia B4:B22
    z["B4"] = kw.id
    z["B5"] = kw.kwota_wnioskowana
    kwota_max = (regulamin.kwota_max_zl if regulamin and regulamin.kwota_max_zl else None) or kw.nabor.kwota_max_zl
    if kwota_max is None:
        kwota_max = kw.kwota_wnioskowana
        uwaga("OSTRZEŻENIE", "Nie znam maksymalnej kwoty naboru – przyjęto kwotę wnioskowaną. Uzupełnij z ogłoszenia.")
    z["B6"] = kwota_max
    z["B7"] = kw.wklad_wlasny.kwota
    z["B8"] = kw.pomysl.data_startu
    forma = kw.forma.opodatkowanie
    if forma == "nie wiem":
        forma = "liniowy"
        uwaga("OSTRZEŻENIE", "Klient nie zna formy opodatkowania – w prognozie przyjęto wariant ostrożny "
                             "(podatek liniowy 19%). Oznacz to w rozmowie z klientem.")
    z["B9"] = forma
    if kw.forma.stawka_ryczaltu:
        z["B10"] = kw.forma.stawka_ryczaltu
    zus = kw.forma.zus
    if zus == "nie wiem":
        zus = "ulga na start + preferencyjny"
        uwaga("OSTRZEŻENIE", "Wariant ZUS nieznany – przyjęto ulgę na start i preferencyjny ZUS. "
                             "Potwierdź, że klient nie prowadził firmy w ostatnich 60 mies.")
    z["B11"] = "TAK" if zus == "ulga na start + preferencyjny" else "NIE"
    z["B13"] = "TAK" if zus in ("ulga na start + preferencyjny", "preferencyjny") else "NIE"
    z["B15"] = "NIE"
    z["B17"] = "TAK" if kw.forma.chorobowe else "NIE"
    if kw.forma.vat == "nie wiem":
        uwaga("OSTRZEŻENIE", "Status VAT nieznany – przyjęto zwolnienie podmiotowe (kwoty brutto).")
    z["B18"] = "TAK" if kw.vat_czynny else "NIE"
    z["B19"] = kw.forma.stawka_vat_sprzedaz
    z["B20"] = kw.prognoza.dni_robocze
    z["B21"] = kw.prognoza.godziny_dzien
    z["B22"] = 0.30

    # --- Kosztorys B..M, wiersze 5–24 (puste wiersze czyścimy)
    for i in range(MAKS_ZAKUPOW):
        r = WIERSZ_ZAKUPOW + i
        zak = kw.zakupy[i] if i < len(kw.zakupy) else None
        ks[f"B{r}"] = zak.nazwa if zak else None
        ks[f"C{r}"] = zak.kategoria if zak else None
        ks[f"D{r}"] = zak.stan if zak else None
        ks[f"E{r}"] = zak.ilosc if zak else None
        ks[f"F{r}"] = zak.cena_netto if zak else None
        ks[f"H{r}"] = zak.stawka_vat if zak else None
        ks[f"L{r}"] = (zak.link or "[brak linku do oferty]") if zak else None
        ks[f"M{r}"] = zak.finansowanie if zak else None
    limity = {
        "sprzęt": 100.0,
        "towar": regulamin.limit_towar_pct if regulamin else None,
        "reklama": regulamin.limit_reklama_pct if regulamin else None,
        "remont": regulamin.limit_remont_pct if regulamin else None,
        "inne": 100.0,
    }
    if regulamin is None:
        uwaga("OSTRZEŻENIE", "Brak regulaminu PUP – limity kategorii w kosztorysie ustawiono na 100% (bez kontroli).")
    for j, kat in enumerate(KATEGORIE):
        lim = limity[kat]
        if lim is None:
            lim = 100.0
            if regulamin is not None:
                uwaga("INFO", f"Regulamin nie podaje limitu dla kategorii „{kat}” – przyjęto 100%.")
        ks[f"B{WIERSZ_LIMITOW + j}"] = lim / 100

    # --- Przychody A6:D10, sezonowość C14:N14, rozruch C19:N19
    for i in range(MAKS_USLUG):
        r = WIERSZ_USLUG + i
        u = kw.uslugi[i] if i < len(kw.uslugi) else None
        pr[f"A{r}"] = u.nazwa if u else None
        pr[f"B{r}"] = u.cena if u else None
        pr[f"C{r}"] = u.liczba_mies if u else None
        pr[f"D{r}"] = u.czas_h if u else None
    for i, kol in enumerate(KOLUMNY_MIESIECY):
        pr[f"{kol}14"] = kw.prognoza.sezonowosc[i]
        pr[f"{kol}19"] = kw.prognoza.rozruch[i]

    # --- Koszty B9:B16, B20
    for j, nazwa in enumerate(KOSZTY_STALE):
        ko[f"B{9 + j}"] = getattr(kw.prognoza.koszty_stale, nazwa)
    ko["B20"] = kw.prognoza.materialy_pct

    plik.parent.mkdir(parents=True, exist_ok=True)
    wb.save(plik)
    return ostrz


# --------------------------------------------------------------------------- przeliczenie
def _soffice() -> str:
    sciezka = os.environ.get("PUP_GEN_SOFFICE") or shutil.which("soffice") or "/usr/bin/soffice"
    if not Path(sciezka).exists():
        raise BladArkusza("Nie znaleziono LibreOffice (soffice). Ustaw PUP_GEN_SOFFICE.")
    return sciezka


def przelicz(plik: Path, docelowy: Path, timeout: int = 180) -> Path:
    """Przelicza formuły: LibreOffice otwiera plik (openpyxl nie zapisuje wartości) i zapisuje jako xlsx."""
    with tempfile.TemporaryDirectory(prefix="pupgen-lo-") as tmp:
        profil = Path(tmp, "profil").as_uri()
        wyjscie = Path(tmp, "out")
        cmd = [_soffice(), f"-env:UserInstallation={profil}", "--headless", "--calc",
               "--convert-to", "xlsx", "--outdir", str(wyjscie), str(plik)]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
        wynik = wyjscie / (plik.stem + ".xlsx")
        if proc.returncode != 0 or not wynik.exists():
            raise BladArkusza(f"LibreOffice nie przeliczył arkusza (kod {proc.returncode}): {proc.stderr[-500:]}")
        docelowy.parent.mkdir(parents=True, exist_ok=True)
        shutil.move(str(wynik), docelowy)
    return docelowy


def konwertuj_pdf(plik: Path, katalog: Path, timeout: int = 180) -> Path:
    with tempfile.TemporaryDirectory(prefix="pupgen-lo-") as tmp:
        profil = Path(tmp, "profil").as_uri()
        cmd = [_soffice(), f"-env:UserInstallation={profil}", "--headless", "--convert-to", "pdf",
               "--outdir", str(katalog), str(plik)]
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=timeout)
    pdf = katalog / (plik.stem + ".pdf")
    if proc.returncode != 0 or not pdf.exists():
        raise BladArkusza(f"Konwersja do PDF nie powiodła się: {proc.stderr[-300:]}")
    return pdf


# --------------------------------------------------------------------------- odczyt
def odczytaj(plik: Path) -> tuple[dict[str, float | str | None], dict[str, float | str | None], int]:
    """Zwraca (wartości ark_*, kontrole Zalozenia!B27:B34, liczba komórek z błędem formuły)."""
    wb = openpyxl.load_workbook(plik, data_only=True)
    wartosci: dict[str, float | str | None] = {}
    for nazwa, dn in wb.defined_names.items():
        if not nazwa.startswith("ark_"):
            continue
        for arkusz, adres in dn.destinations:
            wartosci[nazwa] = wb[arkusz][adres.replace("$", "")].value
    z = wb["Zalozenia"]
    kontrole = {str(z[f"A{r}"].value): z[f"B{r}"].value for r in range(27, 35)}
    bledy = sum(1 for ws in wb.worksheets for row in ws.iter_rows(values_only=True)
                for v in row if isinstance(v, str) and v.startswith("#") and v.rstrip("!?0/").isupper())
    return dict(sorted(wartosci.items())), kontrole, bledy


def odczytaj_tabele(plik: Path) -> dict[str, list[list[str]]]:
    """Tabele z arkusza `Do-wniosku` jako wiersze tekstu (pierwszy wiersz = nagłówek)."""
    wb = openpyxl.load_workbook(plik, data_only=True)
    ws = wb["Do-wniosku"]
    poczatki: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            for klucz, prefiks in (("kalkulacja_miesieczna", "1. "), ("kalkulacja_kwartalna", "1a. "),
                                   ("kosztorys", "2. "), ("wskazniki", "3. ")):
                if a.startswith(prefiks):
                    poczatki[klucz] = r

    def komorka(r: int, c: int) -> str:
        cell = ws.cell(r, c)
        return formatuj_komorke(cell.value, cell.number_format or "")

    tabele: dict[str, list[list[str]]] = {}
    # kalkulacja miesięczna: nagłówek = daty (wiersz +2), dane od +3 do pustego A
    r0 = poczatki["kalkulacja_miesieczna"]
    naglowek = ["Pozycja"] + [komorka(r0 + 2, c) for c in range(3, 15)] + ["Razem 12 mies."]
    wiersze = [naglowek]
    r = r0 + 3
    while ws.cell(r, 1).value:
        wiersze.append([str(ws.cell(r, 1).value)] + [komorka(r, c) for c in range(3, 16)])
        r += 1
    tabele["kalkulacja_miesieczna"] = wiersze
    # kwartalna
    r0 = poczatki["kalkulacja_kwartalna"]
    wiersze = [[str(ws.cell(r0 + 1, c).value or "") for c in (1, 3, 4, 5, 6, 7)]]
    r = r0 + 2
    while ws.cell(r, 1).value and not str(ws.cell(r, 1).value).startswith("Kontrola"):
        wiersze.append([str(ws.cell(r, 1).value)] + [komorka(r, c) for c in range(3, 8)])
        r += 1
    tabele["kalkulacja_kwartalna"] = wiersze
    # specyfikacja wydatków
    r0 = poczatki["kosztorys"]
    wiersze = [[str(ws.cell(r0 + 1, c).value or "") for c in range(1, 11)]]
    r = r0 + 2
    while ws.cell(r, 2).value != "RAZEM" and r < r0 + 40:
        if ws.cell(r, 2).value:
            wiersze.append([komorka(r, c) for c in range(1, 11)])
        r += 1
    for rr in (r, r + 1, r + 2):  # RAZEM, kwota dotacji, wkład własny
        wiersze.append([komorka(rr, c) for c in range(1, 11)])
    tabele["kosztorys"] = wiersze
    # wskaźniki
    r0 = poczatki["wskazniki"]
    wiersze = [["Wskaźnik", "Scenariusz bazowy", "Scenariusz pesymistyczny (−30%)"]]
    r = r0 + 2
    while ws.cell(r, 1).value:
        wiersze.append([str(ws.cell(r, 1).value), komorka(r, 3), komorka(r, 4)])
        r += 1
    tabele["wskazniki"] = wiersze
    return tabele


# --------------------------------------------------------------------------- ostrzeżenia
def ostrzezenia_z_kontroli(wartosci: dict[str, Any], kontrole: dict[str, Any], bledy: int,
                           kw: Kwestionariusz, regulamin: RegulaminPUP | None) -> list[Ostrzezenie]:
    wynik: list[Ostrzezenie] = []

    def dodaj(poziom: str, komunikat: str) -> None:
        wynik.append(Ostrzezenie(poziom=poziom, kategoria="arkusz – kontrole", komunikat=komunikat))

    k = list(kontrole.values())  # kolejność wierszy 27..34
    if bledy:
        dodaj("BLOKUJE", f"Arkusz ma {bledy} komórek z błędem formuły (#VALUE!, #REF! itp.).")
    braki = [n for n, v in wartosci.items() if v is None]
    if braki:
        dodaj("BLOKUJE", f"Puste zmienne arkusza: {', '.join(braki)}.")
    if k[0] != "OK":
        dodaj("BLOKUJE", "Kwota dotacji przekracza maksymalną kwotę naboru (Zalozenia!B27).")
    if k[1] != "OK":
        dodaj("BLOKUJE", "Wydatki z dotacji przekraczają kwotę dotacji (Zalozenia!B28).")
    if isinstance(k[2], (int, float)) and k[2] > 1:
        dodaj("OSTRZEŻENIE", f"Nierozplanowana dotacja: {liczba_pl(k[2], 2)} zł (Zalozenia!B29). "
                             "Komisja oczekuje rozpisania całej kwoty na pozycje.")
    elif isinstance(k[2], (int, float)) and k[2] > 0.005:
        dodaj("INFO", f"Nierozplanowana końcówka dotacji: {liczba_pl(k[2], 2)} zł.")
    if k[3] != "OK":
        dodaj("BLOKUJE", "Przekroczony limit kategorii kosztorysu wg regulaminu (Zalozenia!B30, Kosztorys!E33:E37).")
    stale = wartosci.get("ark_koszty_stale_mies") or 0
    if isinstance(k[4], (int, float)):
        if k[4] < 0:
            dodaj("BLOKUJE", f"Ujemna gotówka w scenariuszu bazowym: najniższe saldo {liczba_pl(k[4])} zł "
                             "(Zalozenia!B31). Zwiększ wkład własny lub obniż koszty.")
        elif k[4] < stale:
            dodaj("OSTRZEŻENIE", f"Niskie saldo gotówki na starcie: {liczba_pl(k[4])} zł, mniej niż miesięczne "
                                 f"koszty stałe ({liczba_pl(stale)} zł). Komisja może dopytać o pierwszy miesiąc.")
    if isinstance(k[5], (int, float)) and k[5] < 0:
        dodaj("OSTRZEŻENIE", f"Ujemna gotówka w scenariuszu pesymistycznym ({liczba_pl(k[5])} zł) – "
                             "opisz plan B w sekcji 09.")
    if isinstance(k[6], (int, float)):
        if k[6] > 1:
            dodaj("BLOKUJE", f"Obłożenie czasu pracy {procent_pl(k[6])} – plan sprzedaży niewykonalny (Zalozenia!B33).")
        elif k[6] > 0.95:
            dodaj("OSTRZEŻENIE", f"Obłożenie czasu pracy {procent_pl(k[6])} – bardzo blisko 100%.")
    if isinstance(k[7], str) and k[7].startswith("UWAGA"):
        dodaj("OSTRZEŻENIE", "Możliwe przekroczenie limitu zwolnienia z VAT (Zalozenia!B34).")
    if isinstance(wartosci.get("ark_pierwszy_mies_zysku"), str):
        dodaj("OSTRZEŻENIE", "Brak miesiąca z zyskiem netto w ciągu 12 miesięcy.")
    zysk = wartosci.get("ark_zysk_netto_12m")
    if isinstance(zysk, (int, float)) and zysk <= 0:
        dodaj("OSTRZEŻENIE", "Zysk netto w 12 miesiącach nie jest dodatni.")
    potrzeby = kw.finanse_domowe.potrzeby_mies
    zysk_pes = wartosci.get("ark_zysk_netto_mies_pes")
    if potrzeby and isinstance(zysk_pes, (int, float)) and zysk_pes < potrzeby:
        dodaj("INFO", f"Średni zysk netto w scenariuszu −30% ({liczba_pl(zysk_pes)} zł) jest niższy niż "
                      f"potrzeby klienta z K11 ({liczba_pl(potrzeby)} zł/mies.).")
    return wynik


# --------------------------------------------------------------------------- całość
def policz(kw: Kwestionariusz, regulamin: RegulaminPUP | None, katalog: Path,
           szablon: Path = PLIK_MODELU) -> WynikFinansowy:
    """Wypełnia, przelicza i odczytuje arkusz. Wynik: `katalog/arkusz-<id>.xlsx`."""
    katalog.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="pupgen-xlsx-") as tmp:
        wejscie = Path(tmp, f"arkusz-{kw.id}.xlsx")
        ostrz = wypelnij(kw, regulamin, wejscie, szablon)
        plik = przelicz(wejscie, katalog / wejscie.name)
    wartosci, kontrole, bledy = odczytaj(plik)
    ostrz += ostrzezenia_z_kontroli(wartosci, kontrole, bledy, kw, regulamin)
    log.info("Arkusz przeliczony: %d zmiennych ark_*, %d ostrzeżeń", len(wartosci), len(ostrz))
    return WynikFinansowy(
        plik=plik,
        wartosci=wartosci,
        ark={n: formatuj_ark(n, v) for n, v in wartosci.items()},
        kontrole=kontrole,
        ostrzezenia=ostrz,
        tabele=odczytaj_tabele(plik),
    )
