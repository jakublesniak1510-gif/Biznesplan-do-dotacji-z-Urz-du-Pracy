"""Składanie DOCX dla klienta (python-docx) i wersji DOCX raportu kontroli dla właściciela."""

from __future__ import annotations

import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_COLOR_INDEX, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .finanse import liczba_pl
from .schemat import Kwestionariusz, MapaFormularza, WynikGeneracji

_RE_BRAK = re.compile(r"(\[BRAK DANYCH:[^\]]*\])")
TYTULY_TABEL = {
    "kosztorys": "Specyfikacja wydatków (arkusz Do-wniosku, część 2)",
    "kalkulacja_miesieczna": "Kalkulacja przychodów i kosztów – 12 miesięcy (arkusz Do-wniosku, część 1)",
    "kalkulacja_kwartalna": "Kalkulacja przychodów i kosztów – kwartały (arkusz Do-wniosku, część 1a)",
    "wskazniki": "Wskaźniki ekonomiczne (arkusz Do-wniosku, część 3)",
}
CO_WPISAC_SAMEMU = [
    "PESEL, seria i numer dowodu osobistego",
    "adres zamieszkania i adres do doręczeń",
    "numer rachunku bankowego",
    "dane poręczycieli albo innej formy zabezpieczenia zwrotu (imiona, PESEL, dochody, dokumenty)",
    "oświadczenia (m.in. o spełnianiu warunków z art. 148 ustawy) – czytasz i podpisujesz sam(a)",
    "formularz informacji o pomocy de minimis i zaświadczenia de minimis",
]


def slug(tekst: str) -> str:
    t = unicodedata.normalize("NFKD", tekst.replace("ł", "l").replace("Ł", "L"))
    t = "".join(c for c in t if not unicodedata.combining(c))
    return re.sub(r"[^A-Za-z0-9]+", "-", t).strip("-")


def nazwa_pliku(wynik: WynikGeneracji, miasto: str, wersja: int = 1) -> str:
    """PUP-[miasto]_[ID]_wniosek_v1.docx (docs/checklista-weryfikacji.md, część H)."""
    return f"PUP-{slug(miasto)}_{wynik.id_zlecenia}_wniosek_v{wersja}.docx"


def _styl(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    for sekcja in doc.sections:
        sekcja.left_margin = sekcja.right_margin = Cm(2)
        sekcja.top_margin = sekcja.bottom_margin = Cm(2)


def _akapit_z_brakami(doc: Document, tekst: str, styl: str | None = None) -> None:
    """Akapit; znaczniki [BRAK DANYCH] podświetlone na żółto, żeby właściciel ich nie przeoczył."""
    p = doc.add_paragraph(style=styl)
    for kawalek in _RE_BRAK.split(tekst):
        if not kawalek:
            continue
        run = p.add_run(kawalek)
        if _RE_BRAK.fullmatch(kawalek):
            run.font.highlight_color = WD_COLOR_INDEX.YELLOW
            run.bold = True


def _powtarzaj_naglowek(wiersz) -> None:
    tr_pr = wiersz._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def _tabela(doc: Document, wiersze: list[list[str]], rozmiar: float = 8) -> None:
    if not wiersze:
        return
    t = doc.add_table(rows=len(wiersze), cols=len(wiersze[0]))
    t.style = "Table Grid"
    for i, wiersz in enumerate(wiersze):
        for j, wartosc in enumerate(wiersz):
            komorka = t.cell(i, j)
            komorka.text = ""
            run = komorka.paragraphs[0].add_run(wartosc)
            run.font.size = Pt(rozmiar)
            run.bold = i == 0 or (j == 0 and wartosc.isupper())
    _powtarzaj_naglowek(t.rows[0])


def _sekcja_pozioma(doc: Document, pozioma: bool) -> None:
    s = doc.add_section()
    if pozioma != (s.orientation == WD_ORIENT.LANDSCAPE):
        s.orientation = WD_ORIENT.LANDSCAPE if pozioma else WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = s.page_height, s.page_width


def _wartosc_danych(klucz: str, wynik: WynikGeneracji, kw: Kwestionariusz) -> str:
    match klucz:
        case "kwota_dotacji":
            return wynik.ark.get("ark_dotacja", f"{liczba_pl(kw.kwota_wnioskowana)} zł")
        case "wklad_wlasny":
            return wynik.ark.get("ark_wklad_wlasny", f"{liczba_pl(kw.wklad_wlasny.kwota)} zł")
        case "forma_opodatkowania":
            f = kw.forma
            if f.opodatkowanie == "ryczałt":
                return f"ryczałt od przychodów ewidencjonowanych ({liczba_pl((f.stawka_ryczaltu or 0) * 100, 1)}%)"
            return {"skala": "zasady ogólne (skala podatkowa)", "liniowy": "podatek liniowy",
                    "nie wiem": "[BRAK DANYCH: forma opodatkowania]"}[f.opodatkowanie]
        case "pkd":
            return "; ".join(f"{p.kod} {p.nazwa}" for p in kw.pomysl.pkd) or "[BRAK DANYCH: kod PKD 2025]"
        case "data_startu":
            return kw.pomysl.data_startu.strftime("%d.%m.%Y")
    return ""


def _strona_tytulowa(doc: Document, wynik: WynikGeneracji, mapa: MapaFormularza, kw: Kwestionariusz) -> None:
    doc.add_heading("Treści do wniosku o dofinansowanie podjęcia działalności gospodarczej", level=0)
    doc.add_paragraph(f"{wynik.pup} · {mapa.formularz}")
    doc.add_paragraph(f"Przygotowano dla: {wynik.imie} ({wynik.miejscowosc}) · data: {wynik.data:%d.%m.%Y}")
    if kw.nabor.termin_zlozenia:
        doc.add_paragraph(f"Termin złożenia wniosku według ogłoszenia: {kw.nabor.termin_zlozenia:%d.%m.%Y}")

    doc.add_heading("Jak korzystać z tego dokumentu", level=1)
    for zdanie in (
        "Treści są ułożone w kolejności pól formularza Twojego urzędu, z nazwami pól jak w formularzu. "
        "Przepisz albo wklej każdą treść do właściwego pola.",
        "Przy każdym polu podajemy limit znaków (ze spacjami). Treści mieszczą się w limitach – po własnych "
        "poprawkach sprawdź długość (Word: Recenzja → Statystyka wyrazów).",
        "Przeczytaj wszystko na głos. Jeśli coś brzmi nie tak, jak powiedział(a)byś to sam(a), popraw – "
        "na rozmowie w urzędzie komisja pyta o treść wniosku.",
        "Tabele liczbowe pochodzą z arkusza finansowego (plik XLSX w paczce). Nie zmieniaj liczb ręcznie – "
        "przy każdej zmianie napisz do nas, przeliczymy arkusz.",
    ):
        doc.add_paragraph(zdanie, style="List Bullet")

    doc.add_heading("Co wpisujesz i podpisujesz sam(a)", level=1)
    for poz in CO_WPISAC_SAMEMU + [f"część formularza: {x}" for x in mapa.poza_zakresem]:
        doc.add_paragraph(poz, style="List Bullet")

    doc.add_heading("Załączniki do przygotowania", level=1)
    zalaczniki = ["oferty (wydruk lub zrzut ekranu z ceną) dla każdej pozycji kosztorysu"]
    if kw.lokalizacja.tryb in ("wynajem", "użyczenie", "własny lokal", "w mieszkaniu"):
        zalaczniki.append("dokument do lokalu (akt własności, umowa najmu lub użyczenia, list intencyjny wynajmującego)")
    zalaczniki += kw.kwalifikacje.dokumenty
    zalaczniki += [f"list intencyjny: {li.od_kogo}" for li in kw.listy_intencyjne]
    for poz in zalaczniki:
        doc.add_paragraph(poz, style="List Bullet")

    doc.add_heading("Ważne", level=1)
    for poz in (
        "Nie rejestruj firmy w CEIDG i nie kupuj niczego z kosztorysu przed podpisaniem umowy z urzędem – "
        "wcześniejsze zakupy nie zostaną rozliczone, a wcześniejsza rejestracja oznacza utratę dotacji.",
        "Złóż wniosek w sposób podany w ogłoszeniu (osobiście, pocztą, przez praca.gov.pl) i zachowaj potwierdzenie.",
        "Jeśli urząd wezwie do uzupełnień, termin jest w piśmie (zwykle 7–14 dni) – od razu prześlij nam pismo.",
        "Wniosek podpisujesz Ty. Usługa nie gwarantuje przyznania dotacji – decyzję podejmuje urząd.",
    ):
        doc.add_paragraph(poz, style="List Bullet")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def zapisz_dokument(wynik: WynikGeneracji, mapa: MapaFormularza, tabele: dict[str, list[list[str]]],
                    kw: Kwestionariusz, plik: Path) -> Path:
    doc = Document()
    _styl(doc)
    _strona_tytulowa(doc, wynik, mapa, kw)
    pola = {p.id: p for p in wynik.pola}
    for pole in mapa.pola:
        wp = pola.get(pole.id)
        doc.add_heading(f"{pole.id}. {pole.nazwa_w_formularzu}", level=1)
        info = []
        if wp and wp.limit_znakow and wp.tekst:
            info.append(f"limit: {wp.limit_znakow} znaków ze spacjami · treść: {wp.liczba_znakow} znaków")
        if pole.uwagi:
            info.append(f"uwagi: {pole.uwagi}")
        if info:
            run = doc.add_paragraph().add_run(" · ".join(info))
            run.italic = True
            run.font.size = Pt(9)
        for klucz in pole.dane:
            doc.add_paragraph(_wartosc_danych(klucz, wynik, kw))
        if wp and wp.tekst:
            for akapit in wp.tekst.split("\n"):
                if akapit.strip():
                    _akapit_z_brakami(doc, akapit.strip())
        for nazwa in pole.tabele:
            pozioma = nazwa == "kalkulacja_miesieczna"
            if pozioma:
                _sekcja_pozioma(doc, True)
            run = doc.add_paragraph().add_run(TYTULY_TABEL[nazwa])
            run.bold = True
            _tabela(doc, tabele.get(nazwa, []), 7 if pozioma else 8)
            if pozioma:
                _sekcja_pozioma(doc, False)
        if not (pole.dane or pole.tabele or (wp and wp.tekst)):
            _akapit_z_brakami(doc, "[BRAK DANYCH: pole do uzupełnienia]")
    plik.parent.mkdir(parents=True, exist_ok=True)
    doc.save(plik)
    return plik


def zapisz_raport_docx(markdown: str, plik: Path) -> Path:
    """Prosta konwersja raportu Markdown (nagłówki, listy, tabela) do DOCX."""
    doc = Document()
    _styl(doc)
    tabela: list[list[str]] = []
    for linia in markdown.splitlines() + [""]:
        if linia.startswith("|"):
            if not re.fullmatch(r"\|[-:| ]+\|", linia):
                tabela.append([c.strip() for c in linia.strip("|").split("|")])
            continue
        if tabela:
            _tabela(doc, tabela, 9)
            tabela = []
        czysta = linia.replace("**", "")
        if linia.startswith("# "):
            doc.add_heading(czysta[2:], level=0)
        elif linia.startswith("## "):
            doc.add_heading(czysta[3:], level=1)
        elif linia.startswith("- "):
            _akapit_z_brakami(doc, czysta[2:], "List Bullet")
        elif linia.startswith("> "):
            doc.add_paragraph().add_run(czysta[2:]).bold = True
        elif linia.strip():
            doc.add_paragraph(czysta)
    doc.save(plik)
    return plik
